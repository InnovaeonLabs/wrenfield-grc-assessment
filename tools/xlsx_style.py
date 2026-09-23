"""openpyxl styling helpers shared by every workbook the build produces."""
from __future__ import annotations

import datetime as dt

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

FONT = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="1F4E5A")
INPUT_FILL = PatternFill("solid", fgColor="FFF7D6")
THIN = Side(style="thin", color="C9D3D6")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BAND_FILLS = {
    "Critical": ("7A1020", "FFFFFF"), "High": ("E4572E", "FFFFFF"),
    "Moderate": ("F4B942", "1A1A1A"), "Low": ("5DA271", "FFFFFF"),
}
STATUS_FILLS = {
    "Satisfied": "D9EFD8", "Other Than Satisfied": "F9D6CF", "Closed": "D9EFD8", "Accepted": "D9EFD8",
    "Risk Accepted": "E6E0F3", "Overdue": "F9D6CF", "Yes": "F9D6CF", "Insufficient": "F9D6CF",
    "Rework Required": "FBE3C8", "Delayed": "FBE3C8",
}


def new_wb() -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)
    wb.calculation.fullCalcOnLoad = True  # formulas have no cached values; Excel/Sheets compute on open
    return wb


def readme_sheet(wb, title: str, lines: list[tuple[str, str]], as_of: dt.date):
    ws = wb.create_sheet("README")
    ws["A1"] = title
    ws["A1"].font = Font(name=FONT, bold=True, size=14, color="1F4E5A")
    ws["A2"] = "Wrenfield Health (fictional) · simulated NIST SP 800-53 Rev. 5 (Release 5.2.0) assessment · synthetic data"
    ws["A2"].font = Font(name=FONT, italic=True, size=9, color="666666")
    ws["A4"] = "As-of (status) date"
    ws["B4"] = as_of
    ws["B4"].number_format = "yyyy-mm-dd"
    ws["B4"].fill = INPUT_FILL
    ws["C4"] = "<- input: all aging/overdue formulas reference this cell (use =TODAY() in a live workbook)"
    for c in ("A4", "C4"):
        ws[c].font = Font(name=FONT, size=10, bold=(c == "A4"))
    ws["B4"].font = Font(name=FONT, size=10, color="0000FF")
    r = 6
    for k, v in lines:
        ws.cell(row=r, column=1, value=k).font = Font(name=FONT, bold=True, size=10)
        c = ws.cell(row=r, column=2, value=v)
        c.font = Font(name=FONT, size=10)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        r += 1
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 110
    return ws


def table_sheet(wb, name: str, headers: list[str], rows: list[list], widths: dict | None = None,
                wrap_cols: set | None = None, date_cols: set | None = None, pct_cols: set | None = None):
    ws = wb.create_sheet(name)
    ws.append(headers)
    for row in rows:
        ws.append(row)
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=i)
        c.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        c.fill = HEADER_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = BORDER
        width = (widths or {}).get(h, 16 if len(h) < 16 else min(len(h) + 2, 30))
        ws.column_dimensions[get_column_letter(i)].width = width
    wrap_cols = wrap_cols or set()
    date_cols = date_cols or set()
    pct_cols = pct_cols or set()
    for r in range(2, ws.max_row + 1):
        for i, h in enumerate(headers, start=1):
            c = ws.cell(row=r, column=i)
            c.font = Font(name=FONT, size=9)
            c.border = BORDER
            c.alignment = Alignment(wrap_text=h in wrap_cols, vertical="top")
            if h in date_cols:
                c.number_format = "yyyy-mm-dd"
            if h in pct_cols:
                c.number_format = "0%"
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"
    ws.row_dimensions[1].height = 32
    return ws


def col_letter(headers: list[str], name: str) -> str:
    return get_column_letter(headers.index(name) + 1)


def add_list_validation(ws, headers: list[str], col_name: str, values: list[str]):
    letter = col_letter(headers, col_name)
    dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True,
                        showErrorMessage=True, errorTitle="Controlled value",
                        error="Choose a value from the list (controlled vocabulary).")
    ws.add_data_validation(dv)
    dv.add(f"{letter}2:{letter}{max(ws.max_row, 2) + 200}")


def band_formatting(ws, headers: list[str], col_name: str):
    letter = col_letter(headers, col_name)
    rng = f"{letter}2:{letter}{ws.max_row + 200}"
    for band, (bg, fg) in BAND_FILLS.items():
        ws.conditional_formatting.add(rng, FormulaRule(
            formula=[f'{letter}2="{band}"'], fill=PatternFill("solid", fgColor=bg), font=Font(name=FONT, color=fg, bold=True)))


def status_formatting(ws, headers: list[str], col_name: str):
    letter = col_letter(headers, col_name)
    rng = f"{letter}2:{letter}{ws.max_row + 200}"
    for val, bg in STATUS_FILLS.items():
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f'{letter}2="{val}"'], fill=PatternFill("solid", fgColor=bg)))
