"""Evaluate every formula in the generated workbooks with pycel (no Excel/LibreOffice needed).

Fails if any formula errors, and cross-checks the combined workbook's Dashboard against the
Python-computed metrics in dashboard/metrics.json, so the spreadsheet and the reports cannot disagree.

    python tools/check_formulas.py
"""
from __future__ import annotations

import json
import sys
import warnings

from openpyxl import load_workbook
from pycel import ExcelCompiler

from common import ROOT

warnings.filterwarnings("ignore")
BOOKS = ["controls/nist-800-53-control-matrix.xlsx", "risk/risk-register.xlsx", "audit/evidence-tracker.xlsx",
         "findings/findings-register.xlsx", "poam/poam.xlsx", "vendor-risk/vendor-assessment.xlsx",
         "workbook/wrenfield-grc-workbook.xlsx"]
ERRORS = ("#NAME?", "#VALUE!", "#REF!", "#DIV/0!", "#N/A", "#NUM!", "#NULL!")


def formula_cells(path):
    wb = load_workbook(path)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    yield ws.title, c.coordinate


def check_book(rel):
    path = ROOT / rel
    xl = ExcelCompiler(filename=str(path))
    n, bad = 0, []
    for sheet, coord in formula_cells(path):
        n += 1
        try:
            v = xl.evaluate(f"'{sheet}'!{coord}")
        except Exception as e:  # noqa: BLE001
            bad.append((sheet, coord, f"EXC {type(e).__name__}: {str(e)[:80]}"))
            continue
        if isinstance(v, str) and v in ERRORS:
            bad.append((sheet, coord, v))
    return xl, n, bad


def main():
    total, failures = 0, []
    dash_vals = {}
    for rel in BOOKS:
        xl, n, bad = check_book(rel)
        total += n
        failures += [(rel, *b) for b in bad]
        print(f"{rel:48} formulas={n:5} errors={len(bad)}")
        if rel.startswith("workbook/"):
            wb = load_workbook(ROOT / rel)
            ws = wb["Dashboard"]
            for r in range(5, ws.max_row + 1):
                mid = ws.cell(row=r, column=1).value
                if mid:
                    dash_vals[mid] = xl.evaluate(f"'Dashboard'!C{r}")
    metrics = json.loads((ROOT / "dashboard" / "metrics.json").read_text(encoding="utf-8"))
    expected = metrics["dashboard_crosscheck"]
    mismatches = []
    for mid, exp in expected.items():
        got = dash_vals.get(mid)
        if got is None or abs(float(got) - float(exp)) > 0.011:
            mismatches.append((mid, exp, got))
    print(f"TOTAL formulas evaluated: {total}; errors: {len(failures)}; dashboard cross-check mismatches: {len(mismatches)}")
    for f in failures[:40]:
        print("  ERROR", f)
    for m in mismatches:
        print("  MISMATCH", m)
    (ROOT / "dashboard" / "formula-check.json").write_text(json.dumps(
        {"formulas_evaluated": total, "errors": len(failures), "dashboard_mismatches": len(mismatches),
         "dashboard_values": {k: (round(v, 4) if isinstance(v, float) else v) for k, v in dash_vals.items()}},
        indent=2) + "\n", encoding="utf-8")
    sys.exit(1 if failures or mismatches else 0)


if __name__ == "__main__":
    main()
