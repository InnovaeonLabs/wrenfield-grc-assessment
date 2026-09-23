"""Build every artifact from the single source of truth in data/.

    python tools/build.py

Outputs (all regenerated; do not hand-edit generated files):
  controls/  nist-800-53-control-matrix.{xlsx,csv}, assessment-procedures.{xlsx,csv}, control-matrix.md
  risk/      risk-register.{xlsx,csv,md}
  access-review/ access-review.xlsx, findings.md          (CSV comes from review_engine.py)
  audit/     evidence-tracker.{xlsx,csv}, evidence-request-list.md
  vendor-risk/ vendor-questionnaire.xlsx, vendor-assessment.{xlsx,md}
  findings/  findings-register.{xlsx,csv,md}
  poam/      poam.{xlsx,csv,md}, lifecycle-traces.md
  docs/      assessment/traceability-matrix.md, assessment/worksheets/*.md, scope/business-requirements.md
  templates/ grc-templates.xlsx + CSV templates
  workbook/  wrenfield-grc-workbook.xlsx  (all sheets + live-formula Dashboard)
"""
from __future__ import annotations

import datetime as dt
from copy import copy
from collections import Counter, defaultdict
from pathlib import Path

from common import META, ROOT, STATUS_DATE, band, d, load_yaml, md_table, read_csv, split_ids, write_csv, write_json
from xlsx_style import (add_list_validation, band_formatting, col_letter, new_wb, readme_sheet, status_formatting,
                        table_sheet)

GEN = "> Generated from `data/*.yaml` by `tools/build.py`. Edit the data, not this file."
ASOF = "README!$B$4"
FAMILY = {
    "AC": "Access Control", "AT": "Awareness and Training", "AU": "Audit and Accountability",
    "CA": "Assessment, Authorization, and Monitoring", "CM": "Configuration Management", "CP": "Contingency Planning",
    "IA": "Identification and Authentication", "IR": "Incident Response", "MP": "Media Protection",
    "PE": "Physical and Environmental Protection", "PL": "Planning", "PM": "Program Management",
    "PS": "Personnel Security", "RA": "Risk Assessment", "SA": "System and Services Acquisition",
    "SC": "System and Communications Protection", "SI": "System and Information Integrity",
    "SR": "Supply Chain Risk Management",
}
VOCAB = {
    "impl": ["Implemented", "Partially Implemented", "Planned", "Not Implemented", "Not Applicable"],
    "result": ["Satisfied", "Other Than Satisfied", "Not Assessed"],
    "level": ["Design", "Implementation", "Operating Effectiveness"],
    "evidence_status": ["Not Requested", "Requested", "Received", "Under Review", "Accepted", "Insufficient",
                        "Rework Required", "Closed"],
    "poam_status": ["Open", "In Progress", "Delayed", "Completed – Pending Validation", "Closed", "Risk Accepted"],
    "severity": ["Critical", "High", "Moderate", "Low"],
    "treatment": ["Mitigate", "Accept", "Transfer", "Avoid"],
    "decision": ["Retain", "Retain (time-bound)", "Retain with exception", "Modify", "Revoke"],
    "vq_eval": ["Satisfactory", "Partially Satisfactory", "Unsatisfactory", "Not Applicable"],
}


def slug(control_id: str) -> str:
    return control_id.replace("(", "-").replace(")", "")


def rating_formula(cell: str) -> str:
    return f'=IF({cell}="","",IF({cell}>=20,"Critical",IF({cell}>=10,"High",IF({cell}>=5,"Moderate","Low"))))'


# ------------------------------------------------------------------ load + derive
def load_all():
    controls = load_yaml("controls.yaml")["controls"]
    fy = load_yaml("findings.yaml")
    D = {
        "controls": controls,
        "ctl": {c["id"]: c for c in controls},
        "reqs": {r["id"]: r for r in load_yaml("requirements.yaml")},
        "evidence": load_yaml("evidence.yaml")["evidence"],
        "findings": fy["findings"],
        "obs": {o["id"]: o for o in fy["observations"]},
        "risks": load_yaml("risks.yaml")["risks"],
        "poam": load_yaml("poam.yaml")["poam"],
        "vendor": load_yaml("vendor.yaml"),
        "worksheets": load_yaml("worksheets.yaml")["worksheets"],
        "ar": read_csv(ROOT / "access-review" / "access-review.csv"),
    }
    D["ev"] = {e["id"]: e for e in D["evidence"]}
    D["fnd"] = {f["id"]: f for f in D["findings"]}
    D["rsk"] = {r["id"]: r for r in D["risks"]}
    D["pm"] = {p["id"]: p for p in D["poam"]}
    D["ws"] = {w["control"]: w for w in D["worksheets"]}
    # derived: finding severity
    for f in D["findings"]:
        f["score"] = f["likelihood"] * f["impact"]
        f["severity"] = band(f["score"])
    # derived: risk scores
    for r in D["risks"]:
        for k in ("inherent", "current", "residual"):
            r[f"{k}_score"] = r[k][0] * r[k][1]
            r[f"{k}_band"] = band(r[f"{k}_score"])
    # derived: POA&M
    sla = META["poam_sla_days"]
    for p in D["poam"]:
        f = D["fnd"][p["finding"]]
        p["severity"] = f["severity"]
        p["controls"] = f["controls"]
        p["risks"] = f["risks"]
        p["open"] = d(p["open_date"])
        p["sla_date"] = p["open"] + dt.timedelta(days=sla[f["severity"]])
        p["target"] = d(p["current_target"])
        p["closed"] = d(p.get("closure_date"))
        done = sum(m["status"] == "Done" for m in p["milestones"])
        p["pct"] = done / len(p["milestones"])
        p["days_open"] = ((p["closed"] or STATUS_DATE) - p["open"]).days
        p["overdue"] = p["status"] not in ("Closed", "Risk Accepted") and p["target"] < STATUS_DATE
        p["slipped"] = str(p["original_target"]) != str(p["current_target"])
        p["within_sla"] = p["status"] == "Risk Accepted" or p["target"] <= p["sla_date"]  # accepted: target = acceptance expiry
        p["late_milestones"] = [m["id"] for m in p["milestones"]
                                if m["status"] != "Done" and d(m["due"]) < STATUS_DATE]
    # derived: evidence
    for e in D["evidence"]:
        e["in_repo"] = "Yes" if (ROOT / str(e["location"])).exists() else "No"
        e["sufficiency"] = {"Accepted": "Sufficient", "Closed": "Sufficient", "Insufficient": "Insufficient",
                            "Rework Required": "Insufficient (rework)"}.get(e["status"], "Pending")
        rec, due = d(e.get("received")), d(e["due_date"])
        e["days_late"] = max(0, (rec - due).days) if rec else max(0, (STATUS_DATE - due).days)
    return D


# ------------------------------------------------------------------ controls
CTRL_HEADERS = ["Control ID", "Control Family", "Control Name", "Baseline (SP 800-53B)", "Control Requirement Summary",
                "Applicability", "Applicability Rationale", "Control Origination", "System / Process", "Control Owner",
                "Business Requirements (REQ)", "Implementation Status (Owner-Stated)",
                "Implementation Status (Assessor-Validated)", "Status Changed by Assessor?",
                "Implementation Description", "Evidence Required", "Evidence Available (EVID)", "Evidence Location",
                "Highest Evidence Level", "Testing Method", "Test Procedure / Expected vs Actual", "Assessment Result",
                "Gap Identified", "Risk Rating", "Related Risk ID", "Related Finding ID", "Related POA&M ID",
                "POA&M Status (as of status date)", "Observations", "Last Reviewed", "Next Review",
                "CSF 2.0 Function (informal)", "Notes"]


def control_rows(D):
    rows = []
    for c in D["controls"]:
        pstat = "; ".join(f"{p}: {D['pm'][p]['status']}" for p in c["poams"]) or ("n/a" if c["result"] != "Other Than Satisfied" else "")
        locs = "; ".join(str(D["ev"][e]["location"]) for e in c["evidence"])
        rows.append({
            "Control ID": c["id"], "Control Family": f"{c['family']} - {FAMILY[c['family']]}", "Control Name": c["name"],
            "Baseline (SP 800-53B)": c["baseline"], "Control Requirement Summary": c["requirement"],
            "Applicability": c["applicability"], "Applicability Rationale": c.get("applicability_rationale", ""),
            "Control Origination": c["origination"], "System / Process": c["systems"], "Control Owner": c["owner"],
            "Business Requirements (REQ)": "; ".join(c["requirements"]),
            "Implementation Status (Owner-Stated)": c["status_owner"],
            "Implementation Status (Assessor-Validated)": c["status_validated"],
            "Status Changed by Assessor?": "Yes" if c["status_owner"] != c["status_validated"] else "No",
            "Implementation Description": c["implementation"], "Evidence Required": c["evidence_required"],
            "Evidence Available (EVID)": "; ".join(c["evidence"]), "Evidence Location": locs,
            "Highest Evidence Level": c["evidence_level"], "Testing Method": ", ".join(c["methods"]),
            "Test Procedure / Expected vs Actual": c["test_summary"], "Assessment Result": c["result"],
            "Gap Identified": c["gap"], "Risk Rating": c["risk_rating"], "Related Risk ID": "; ".join(c["risks"]),
            "Related Finding ID": "; ".join(c["findings"]), "Related POA&M ID": "; ".join(c["poams"]),
            "POA&M Status (as of status date)": pstat, "Observations": "; ".join(c["obs"]),
            "Last Reviewed": d(c["last_reviewed"]), "Next Review": d(c["next_review"]),
            "CSF 2.0 Function (informal)": c["csf"], "Notes": c.get("notes", ""),
        })
    return rows


def sheet_controls(wb, D):
    rows = control_rows(D)
    widths = {"Control ID": 11, "Control Name": 26, "Control Requirement Summary": 50, "Implementation Description": 60,
              "Test Procedure / Expected vs Actual": 70, "Gap Identified": 36, "Evidence Location": 40,
              "Applicability Rationale": 30, "Notes": 36, "Evidence Required": 36}
    wrap = {"Control Requirement Summary", "Implementation Description", "Test Procedure / Expected vs Actual",
            "Gap Identified", "Evidence Location", "Applicability Rationale", "Notes", "Evidence Required", "Control Name"}
    ws = table_sheet(wb, "Controls", CTRL_HEADERS, [[r[h] for h in CTRL_HEADERS] for r in rows], widths, wrap,
                     {"Last Reviewed", "Next Review"})
    o, v = col_letter(CTRL_HEADERS, "Implementation Status (Owner-Stated)"), col_letter(CTRL_HEADERS, "Implementation Status (Assessor-Validated)")
    chg = col_letter(CTRL_HEADERS, "Status Changed by Assessor?")
    for r in range(2, ws.max_row + 1):
        ws[f"{chg}{r}"] = f'=IF({o}{r}<>{v}{r},"Yes","No")'
    add_list_validation(ws, CTRL_HEADERS, "Implementation Status (Owner-Stated)", VOCAB["impl"])
    add_list_validation(ws, CTRL_HEADERS, "Implementation Status (Assessor-Validated)", VOCAB["impl"])
    add_list_validation(ws, CTRL_HEADERS, "Assessment Result", VOCAB["result"])
    add_list_validation(ws, CTRL_HEADERS, "Highest Evidence Level", VOCAB["level"])
    band_formatting(ws, CTRL_HEADERS, "Risk Rating")
    status_formatting(ws, CTRL_HEADERS, "Assessment Result")
    status_formatting(ws, CTRL_HEADERS, "Status Changed by Assessor?")
    return rows


PROC_HEADERS = ["Test Procedure ID", "Control ID", "Control Name", "Method", "Assessment Object / Procedure",
                "Expected Result", "Actual Result", "Step Result", "Evidence", "Control Determination",
                "Deficiency", "Related Finding", "Worksheet"]


def procedure_rows(D):
    rows = []
    for w in D["worksheets"]:
        c = D["ctl"][w["control"]]
        for s in w["steps"]:
            rows.append({"Test Procedure ID": s["tp"], "Control ID": w["control"], "Control Name": c["name"],
                         "Method": s["method"], "Assessment Object / Procedure": s["procedure"],
                         "Expected Result": s["expected"], "Actual Result": s["actual"], "Step Result": s["result"],
                         "Evidence": s["evidence"], "Control Determination": w["conclusion"],
                         "Deficiency": w["deficiency"], "Related Finding": "; ".join(w["findings"]),
                         "Worksheet": f"docs/assessment/worksheets/{slug(w['control'])}.md"})
    return rows


def build_controls(D):
    wb = new_wb()
    readme_sheet(wb, "NIST SP 800-53 Rev. 5 Control Matrix: WCP-PROD", [
        ("Catalog", META["catalog"] + ", verified against NIST CSRC and OSCAL content on 2026-09-23"),
        ("Baseline", "SP 800-53B Moderate (OSCAL profile v5.2.0), tailored to 61 controls. PM-9 is organization-level (PM family is not allocated to baselines)"),
        ("Two status columns", "Owner-Stated = what the SSP/control owner claimed. Assessor-Validated = determination after testing. 'Status Changed by Assessor?' is a live formula"),
        ("Assessment Result", "SP 800-53A terms: Satisfied / Other Than Satisfied / Not Assessed (N/A controls only)"),
        ("Evidence level", "Design (documentation exists) < Implementation (configured/deployed) < Operating Effectiveness (worked over the period)"),
        ("Controlled values", "Status columns use dropdown validation (see Lists in the combined workbook)"),
        ("Traceability", "REQ -> control -> EVID -> test (TP-) -> result -> RISK -> FIND -> POAM. See docs/assessment/traceability-matrix.md"),
        ("Source of truth", "data/controls.yaml, generated by tools/build.py"),
    ], STATUS_DATE)
    rows = sheet_controls(wb, D)
    prows = procedure_rows(D)
    wb.save(ROOT / "controls" / "nist-800-53-control-matrix.xlsx")
    write_csv(ROOT / "controls" / "nist-800-53-control-matrix.csv", rows, CTRL_HEADERS)

    wb2 = new_wb()
    readme_sheet(wb2, "Assessment Procedures (SP 800-53A-style)", [
        ("Purpose", "Examine / Interview / Test steps for 15 key controls with expected vs actual results"),
        ("Objectives", "Paraphrased from SP 800-53A Rev. 5 assessment objectives; the NIST publication is authoritative"),
        ("Step Result", "Pass / Fail / Info. A control is Satisfied only if no step fails"),
        ("Worksheets", "Narrative worksheets: docs/assessment/worksheets/"),
    ], STATUS_DATE)
    ws = table_sheet(wb2, "Procedures", PROC_HEADERS, [[r[h] for h in PROC_HEADERS] for r in prows],
                     {"Assessment Object / Procedure": 50, "Expected Result": 34, "Actual Result": 50, "Deficiency": 34,
                      "Control Determination": 22},
                     {"Assessment Object / Procedure", "Expected Result", "Actual Result", "Deficiency", "Control Determination"})
    status_formatting(ws, PROC_HEADERS, "Step Result")
    wb2.save(ROOT / "controls" / "assessment-procedures.xlsx")
    write_csv(ROOT / "controls" / "assessment-procedures.csv", prows, PROC_HEADERS)

    # markdown summary
    fam = defaultdict(Counter)
    for c in D["controls"]:
        fam[c["family"]][c["result"]] += 1
    lines = [f"# Control Matrix Summary: NIST SP 800-53 Rev. 5 (Release 5.2.0)", "", GEN, "",
             "Full matrix (33 fields): [nist-800-53-control-matrix.xlsx](nist-800-53-control-matrix.xlsx) · "
             "[CSV](nist-800-53-control-matrix.csv) · Procedures: [assessment-procedures.csv](assessment-procedures.csv)", "",
             "## Results by family", "",
             md_table([{"Family": f"{k} {FAMILY[k]}", "Controls": sum(v.values()), "Satisfied": v["Satisfied"],
                        "Other Than Satisfied": v["Other Than Satisfied"], "N/A": v["Not Assessed"]}
                       for k, v in sorted(fam.items())], ["Family", "Controls", "Satisfied", "Other Than Satisfied", "N/A"]),
             "", "## All controls", "",
             md_table([{"ID": c["id"], "Name": c["name"], "Owner-stated": c["status_owner"],
                        "Validated": c["status_validated"], "Result": c["result"], "Evidence level": c["evidence_level"],
                        "Risk": c["risk_rating"], "Finding": ", ".join(c["findings"]) or "—",
                        "POA&M": ", ".join(c["poams"]) or "—",
                        "Worksheet": f"[{c['id']}](../docs/assessment/worksheets/{slug(c['id'])}.md)" if c["id"] in D["ws"] else ""}
                       for c in D["controls"]],
                      ["ID", "Name", "Owner-stated", "Validated", "Result", "Evidence level", "Risk", "Finding", "POA&M", "Worksheet"])]
    (ROOT / "controls" / "control-matrix.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
    return rows


def build_worksheets_md(D):
    out = ROOT / "docs" / "assessment" / "worksheets"
    out.mkdir(parents=True, exist_ok=True)
    index = []
    for w in D["worksheets"]:
        c = D["ctl"][w["control"]]
        steps = md_table([{"TP": s["tp"], "Method": s["method"], "Procedure": s["procedure"], "Expected": s["expected"],
                           "Actual": s["actual"], "Result": ("✅ " if s["result"] == "Pass" else "❌ " if s["result"] == "Fail" else "ℹ️ ") + s["result"],
                           "Evidence": s["evidence"]} for s in w["steps"]],
                         ["TP", "Method", "Procedure", "Expected", "Actual", "Result", "Evidence"])
        risks = ", ".join(f"{r} ({D['rsk'][r]['title']})" for r in c["risks"]) or "None"
        finds = ", ".join(f"[{f}](../../../findings/findings-register.md#{f.lower()}) {D['fnd'][f]['title']} ({D['fnd'][f]['severity']})"
                          for f in w["findings"]) or "None"
        body = f"""# Control Assessment Worksheet: {c['id']} {c['name']}

{GEN}

| Field | Value |
|---|---|
| Control | **{c['id']} {c['name']}** ({FAMILY[c['family']]}) · baseline: {c['baseline']} |
| Requirement (paraphrased) | {c['requirement']} |
| Business drivers | {', '.join(f"{r} ({D['reqs'][r]['source']})" for r in c['requirements'])} |
| System / process | {c['systems']} |
| Control owner | {c['owner']} |
| Origination | {c['origination']} |
| Owner-stated status → validated | {c['status_owner']} → **{c['status_validated']}** |
| Assessment objective (SP 800-53A, paraphrased) | {w['objective']} |
| Interviewees | {', '.join(w['interviewees'])} |

## Implementation (as described and observed)
{c['implementation']}

## Test steps
{steps}

## Determination
| | |
|---|---|
| **Expected vs actual (summary)** | {c['test_summary']} |
| **Assessment result** | **{w['conclusion']}** |
| **Highest evidence level obtained** | {c['evidence_level']} |
| **Deficiency** | {w['deficiency'] or 'None'} |
| **Risk** | {risks} |
| **Finding** | {finds} |
| **POA&M** | {', '.join(c['poams']) or 'None'} |
| **Control owner response** | {w['owner_response']} |

## Quality review (4 perspectives)
- **Auditor:** every step cites evidence IDs that resolve in the [evidence tracker](../../../audit/evidence-tracker.csv); the conclusion follows from the failed or passed steps.
- **Control owner:** the deficiency names a specific, actionable gap (not "improve access management").
- **Risk manager:** the linked risk states the business consequence and carries a scored rationale.
- **Hiring manager:** shows test design (population, sample, criteria), not just a checklist tick.
"""
        (out / f"{slug(c['id'])}.md").write_text(body, encoding="utf-8")
        index.append({"Control": f"[{c['id']}]({slug(c['id'])}.md)", "Name": c["name"], "Steps": len(w["steps"]),
                      "Result": w["conclusion"], "Finding": ", ".join(w["findings"]) or "—"})
    (out / "README.md").write_text("# Control Assessment Worksheets\n\n" + GEN + "\n\nBlank template: "
                                   "[templates/control-assessment-worksheet.md](../../../templates/control-assessment-worksheet.md)\n\n"
                                   + md_table(index, ["Control", "Name", "Steps", "Result", "Finding"]) + "\n", encoding="utf-8")


def build_traceability(D):
    rows = []
    for c in D["controls"]:
        tps = [s["tp"] for s in D["ws"][c["id"]]["steps"]] if c["id"] in D["ws"] else []
        impl = c["implementation"].split(". ")[0].rstrip(".") + "."
        rem = "; ".join(f"{p} ({D['pm'][p]['status']})" for p in c["poams"]) or ("Monitor" if c["result"] == "Satisfied" else "—")
        rows.append({"Control": c["id"], "Business requirement": ", ".join(c["requirements"]), "Implementation": impl,
                     "Evidence": ", ".join(c["evidence"]),
                     "Test procedure": (", ".join(c["methods"]) + (f" ({tps[0]}…{tps[-1]})" if tps else "")),
                     "Result": ("✅ " if c["result"] == "Satisfied" else "❌ " if c["result"] == "Other Than Satisfied" else "— ") + c["result"],
                     "Risk": ", ".join(c["risks"]) or "—", "Finding": ", ".join(c["findings"]) or "—", "Remediation": rem})
    req_rows = [{"REQ": r["id"], "Source": r["source"], "Requirement": r["requirement"],
                 "Controls": ", ".join(c["id"] for c in D["controls"] if r["id"] in c["requirements"])}
                for r in D["reqs"].values()]
    txt = f"""# Traceability Matrix

{GEN}

Each row answers one auditor question: *how does this control get from contract language to measurable evidence, and what happened when it was tested?*

```
CONTROL → BUSINESS REQUIREMENT → IMPLEMENTATION → EVIDENCE → TEST PROCEDURE → RESULT → RISK → FINDING → REMEDIATION
```

{md_table(rows, list(rows[0].keys()))}

## Requirement coverage (reverse view)
Every business requirement maps to at least one tested control ([tools/validate.py](../../tools/validate.py) enforces this).

{md_table(req_rows, ["REQ", "Source", "Requirement", "Controls"])}
"""
    (ROOT / "docs" / "assessment" / "traceability-matrix.md").write_text(txt, encoding="utf-8")
    # business requirements doc
    (ROOT / "docs" / "scope" / "business-requirements.md").write_text(
        f"# Business Requirements Register\n\n{GEN}\n\nRequirements that justify control selection. SHCA CSA terms and customer "
        f"BAA terms are fictional; HIPAA citations refer to 45 CFR Part 164.\n\n"
        + md_table([{**r, "Type": D["reqs"][r["REQ"]]["type"]} for r in req_rows], ["REQ", "Type", "Source", "Requirement", "Controls"]) + "\n",
        encoding="utf-8")


# ------------------------------------------------------------------ risk
RISK_HEADERS = ["Risk ID", "Risk Title", "Risk Description", "Asset / Process", "Threat", "Vulnerability",
                "Existing Controls", "Inherent Likelihood", "Inherent Impact", "Inherent Score", "Inherent Risk",
                "Current Likelihood", "Current Impact", "Current Score", "Current Risk", "Treatment Strategy",
                "Planned Mitigation", "Risk Owner", "Target Date", "Residual Likelihood", "Residual Impact",
                "Residual Score", "Residual Risk", "Status", "Related NIST Controls", "Related Findings",
                "Related POA&M", "Review Date", "Rationale: Inherent", "Rationale: Current", "Rationale: Residual"]


def risk_rows(D):
    return [{
        "Risk ID": r["id"], "Risk Title": r["title"], "Risk Description": r["description"], "Asset / Process": r["asset"],
        "Threat": r["threat"], "Vulnerability": r["vulnerability"], "Existing Controls": r["existing_controls"],
        "Inherent Likelihood": r["inherent"][0], "Inherent Impact": r["inherent"][1], "Inherent Score": r["inherent_score"],
        "Inherent Risk": r["inherent_band"], "Current Likelihood": r["current"][0], "Current Impact": r["current"][1],
        "Current Score": r["current_score"], "Current Risk": r["current_band"], "Treatment Strategy": r["treatment"],
        "Planned Mitigation": r["mitigation"], "Risk Owner": r["owner"], "Target Date": d(r["target_date"]),
        "Residual Likelihood": r["residual"][0], "Residual Impact": r["residual"][1], "Residual Score": r["residual_score"],
        "Residual Risk": r["residual_band"], "Status": r["status"], "Related NIST Controls": "; ".join(r["controls"]),
        "Related Findings": "; ".join(r["findings"]), "Related POA&M": "; ".join(r["poams"]),
        "Review Date": d(r["review_date"]), "Rationale: Inherent": r["rationale_inherent"],
        "Rationale: Current": r["rationale_current"], "Rationale: Residual": r["rationale_residual"],
    } for r in D["risks"]]


def sheet_risks(wb, D):
    rows = risk_rows(D)
    wrap = {"Risk Title", "Risk Description", "Threat", "Vulnerability", "Existing Controls", "Planned Mitigation",
            "Rationale: Inherent", "Rationale: Current", "Rationale: Residual", "Asset / Process"}
    ws = table_sheet(wb, "Risks", RISK_HEADERS, [[r[h] for h in RISK_HEADERS] for r in rows],
                     {h: 40 for h in wrap} | {"Risk Title": 30}, wrap, {"Target Date", "Review Date"})
    for pre in ("Inherent", "Current", "Residual"):
        L, I = col_letter(RISK_HEADERS, f"{pre} Likelihood"), col_letter(RISK_HEADERS, f"{pre} Impact")
        S, B = col_letter(RISK_HEADERS, f"{pre} Score"), col_letter(RISK_HEADERS, f"{pre} Risk")
        for r in range(2, ws.max_row + 1):
            ws[f"{S}{r}"] = f"={L}{r}*{I}{r}"
            ws[f"{B}{r}"] = rating_formula(f"{S}{r}")
        band_formatting(ws, RISK_HEADERS, f"{pre} Risk")
        for col in (f"{pre} Likelihood", f"{pre} Impact"):
            from openpyxl.worksheet.datavalidation import DataValidation
            dv = DataValidation(type="whole", operator="between", formula1="1", formula2="5", showErrorMessage=True,
                                error="Likelihood and impact use the 1-5 scale in risk/risk-methodology.md")
            ws.add_data_validation(dv)
            letter = col_letter(RISK_HEADERS, col)
            dv.add(f"{letter}2:{letter}{ws.max_row + 200}")
    add_list_validation(ws, RISK_HEADERS, "Treatment Strategy", VOCAB["treatment"])
    return rows


def heatmap(D, state: str) -> str:
    grid = defaultdict(list)
    for r in D["risks"]:
        grid[tuple(r[state])].append(r["id"].replace("RISK-", "R"))
    lines = ["| Likelihood ↓ / Impact → | 1 | 2 | 3 | 4 | 5 |", "|---|---|---|---|---|---|"]
    for L in range(5, 0, -1):
        cells = []
        for I in range(1, 6):
            ids = grid.get((L, I), [])
            cells.append(f"{band(L * I)[0]}: {', '.join(ids)}" if ids else "")
        lines.append(f"| **{L}** | " + " | ".join(cells) + " |")
    return "\n".join(lines)


def build_risk(D):
    wb = new_wb()
    readme_sheet(wb, "Enterprise Cybersecurity Risk Register", [
        ("Method", "RSK-MTH-001 (risk/risk-methodology.md): 1-5 likelihood x 1-5 impact, adapted from NIST SP 800-30 Rev. 1"),
        ("Formulas", "Score = Likelihood x Impact. Rating = IF(score>=20,Critical, >=10 High, >=5 Moderate, else Low). Edit L/I and scores recalculate"),
        ("Three states", "Inherent (no controls) -> Current (controls as tested; failed controls earn no credit) -> Residual (target after treatment)"),
        ("Rationale", "Every L and I value has a written rationale (right-most columns). A score without rationale is not accepted"),
        ("Limitation", "Ordinal scales: multiplication is a prioritization convention, not a measurement. Ties broken by Impact, then Likelihood"),
    ], STATUS_DATE)
    rows = sheet_risks(wb, D)
    wb.save(ROOT / "risk" / "risk-register.xlsx")
    write_csv(ROOT / "risk" / "risk-register.csv", rows, RISK_HEADERS)
    cnt = {s: Counter(r[f"{s}_band"] for r in D["risks"]) for s in ("inherent", "current", "residual")}
    band_table = md_table([{"State": s.title(), **{b: cnt[s][b] for b in VOCAB["severity"]}} for s in cnt],
                          ["State"] + VOCAB["severity"])
    summary = md_table([{"ID": r["id"], "Risk": r["title"], "Inherent": f"{r['inherent_score']} {r['inherent_band']}",
                         "Current": f"{r['current_score']} {r['current_band']}", "Treatment": r["treatment"],
                         "Residual": f"{r['residual_score']} {r['residual_band']}", "Owner": r["owner"].split(" (")[0],
                         "Status": r["status"], "POA&M": ", ".join(r["poams"]) or "—"} for r in D["risks"]],
                       ["ID", "Risk", "Inherent", "Current", "Treatment", "Residual", "Owner", "Status", "POA&M"])
    details = []
    for r in D["risks"]:
        details.append(f"""### {r['id']}: {r['title']}
*{r['description']}*

| Field | Value |
|---|---|
| Asset / process | {r['asset']} |
| Threat → vulnerability | {r['threat']} → {r['vulnerability']} |
| Existing controls | {r['existing_controls']} |
| **Inherent** | L{r['inherent'][0]} × I{r['inherent'][1]} = **{r['inherent_score']} {r['inherent_band']}**: {r['rationale_inherent']} |
| **Current** (as tested) | L{r['current'][0]} × I{r['current'][1]} = **{r['current_score']} {r['current_band']}**: {r['rationale_current']} |
| **Treatment** | {r['treatment']}: {r['mitigation']} |
| **Residual** (target) | L{r['residual'][0]} × I{r['residual'][1]} = **{r['residual_score']} {r['residual_band']}**: {r['rationale_residual']} |
| Owner · target · status | {r['owner']} · {r['target_date']} · {r['status']} |
| Links | Controls {', '.join(r['controls'])} · Findings {', '.join(r['findings']) or '—'} · POA&M {', '.join(r['poams']) or '—'} |
""")
    txt = f"""# Enterprise Cybersecurity Risk Register

{GEN}

Workbook with live formulas: [risk-register.xlsx](risk-register.xlsx) · CSV: [risk-register.csv](risk-register.csv) · Method: [risk-methodology.md](risk-methodology.md) (approved **before** scoring)

Status date {STATUS_DATE}. {len(D['risks'])} risks.

## Rating distribution: inherent → current → residual
{band_table}

No risk reaches zero. Residual ratings reflect the target after treatment. Risks stay open until the POA&M item is validated.

## Heat map: current risk (as tested)
{heatmap(D, 'current')}

## Heat map: residual risk (target)
{heatmap(D, 'residual')}

## Register
{summary}

## Scoring rationale (inherent → control → residual)
{chr(10).join(details)}
"""
    (ROOT / "risk" / "risk-register.md").write_text(txt, encoding="utf-8")
    return rows


# ------------------------------------------------------------------ access review
AR_ORDER = ["Row ID", "User / Account", "Account Type", "Department", "Role", "Manager", "HR Status", "System",
            "Access Level", "Privileged?", "PHI Access?", "Date Granted", "Last Login", "Business Justification",
            "Reviewer", "Review Date", "Engine Flags", "Review Decision", "Exception ID", "Issue Identified",
            "Required Action", "Action Owner", "Due Date", "Completion Status", "Completion Date", "Closure Evidence",
            "Related Finding", "Affected Control", "Reviewer Comment"]


def sheet_access(wb, D):
    date_cols = {"Date Granted", "Last Login", "Review Date", "Due Date", "Completion Date"}
    rows = [[d(r[h]) if h in date_cols else r[h] for h in AR_ORDER] for r in D["ar"]]
    ws = table_sheet(wb, "AccessReview", AR_ORDER, rows,
                     {"Issue Identified": 44, "Required Action": 40, "Reviewer Comment": 40, "Business Justification": 34,
                      "Access Level": 28, "Role": 26},
                     {"Issue Identified", "Required Action", "Reviewer Comment", "Business Justification"}, date_cols)
    add_list_validation(ws, AR_ORDER, "Review Decision", VOCAB["decision"])
    status_formatting(ws, AR_ORDER, "Completion Status")
    # helper: 1 on the first row of each exception ID, so distinct exceptions = SUM(helper)
    hc = len(AR_ORDER) + 1
    exl = col_letter(AR_ORDER, "Exception ID")
    ws.cell(row=1, column=hc, value="First Row of Exception?")
    ws.cell(row=1, column=hc).font = copy(ws.cell(row=1, column=1).font)
    ws.cell(row=1, column=hc).fill = copy(ws.cell(row=1, column=1).fill)
    for r in range(2, ws.max_row + 1):
        ws.cell(row=r, column=hc, value=f'=IF({exl}{r}="",0,IF(COUNTIF({exl}$2:{exl}{r},{exl}{r}&"")=1,1,0))')
    return ws


def build_access_review(D):
    summary = __import__("json").loads((ROOT / "access-review" / "summary.json").read_text(encoding="utf-8"))
    wb = new_wb()
    readme_sheet(wb, "Access-Control Review (Privileged and PHI Access Certification)", [
        ("Snapshot", f"{summary['snapshot_date']}; certification performed 2026-07-06 to 2026-07-17"),
        ("Population", f"{summary['entitlement_rows_reviewed']} entitlements, {summary['identities_reviewed']} accounts, {len(summary['systems'])} system labels; "
                       f"{summary['privileged_rows']} privileged, {summary['phi_rows']} PHI-bearing"),
        ("Method", "tools/review_engine.py reconciles HR, Okta, entitlement exports, service register, role baseline, SoD rules and raises flags; reviewers disposition every flag"),
        ("Result", f"{summary['flagged_rows']} flagged rows -> {summary['exceptions']} exceptions ({summary['exception_rows']} rows) + "
                   f"{summary['false_positive_rows']} false positive dispositioned; {summary['exceptions_closed']} closed, "
                   f"{summary['exceptions_risk_accepted']} risk-accepted, {summary['exceptions_open']} open"),
        ("Rule", "No self-review: reviewers never certify their own access (e.g., CTO reviews the Director of Platform Engineering's admin rows)"),
    ], STATUS_DATE)
    sheet_access(wb, D)
    flag_rows = [[k, v["meaning"], v["count"]] for k, v in summary["flags_by_type"].items()]
    table_sheet(wb, "Flag Legend", ["Flag", "Meaning", "Rows Flagged"], flag_rows, {"Meaning": 70}, {"Meaning"})
    ex_rows = [[k, v["issue"], ", ".join(v["accounts"]), v["rows"], v["finding"], v["status"]] for k, v in summary["exceptions_detail"].items()]
    table_sheet(wb, "Exceptions", ["Exception ID", "Issue", "Accounts", "Rows", "Finding", "Status"], ex_rows,
                {"Issue": 70, "Accounts": 30}, {"Issue"})
    wb.save(ROOT / "access-review" / "access-review.xlsx")

    # findings.md: every exception traced FINDING -> RISK -> CONTROL -> REMEDIATION -> OWNER -> DUE -> CLOSURE EVIDENCE
    by_ex = defaultdict(list)
    for r in D["ar"]:
        if r["Exception ID"]:
            by_ex[r["Exception ID"]].append(r)
    blocks, table = [], []
    for ex in sorted(by_ex):
        rs = by_ex[ex]
        f = D["fnd"][rs[0]["Related Finding"]]
        risks = ", ".join(f"{x} ({D['rsk'][x]['title']})" for x in f["risks"])
        accts = ", ".join(sorted({f"`{r['User / Account']}` @ {r['System']}" for r in rs}))
        status = summary["exceptions_detail"][ex]["status"]
        comp = "; ".join(sorted({r["Completion Date"] for r in rs if r["Completion Date"]})) or "open"
        table.append({"Exception": ex, "Accounts": accts, "Decision": rs[0]["Review Decision"], "Finding": f["id"],
                      "Status": status})
        blocks.append(f"""### {ex}: {rs[0]['Issue Identified']}
| Chain | Detail |
|---|---|
| **Accounts / systems** | {accts} |
| **Engine flags** | {'; '.join(sorted({fl for r in rs for fl in r['Engine Flags'].split('; ') if fl})) or 'Reviewer-identified'} |
| **Finding** | {f['id']}: {f['title']} ({f['severity']}) |
| **Risk** | {risks} |
| **Affected control** | {rs[0]['Affected Control']} |
| **Remediation** | {rs[0]['Review Decision']}: {rs[0]['Required Action']} |
| **Owner** | {rs[0]['Action Owner']} |
| **Due date** | {rs[0]['Due Date']} |
| **Status / completed** | {status} / {comp} |
| **Evidence of closure** | {'; '.join(sorted({r['Closure Evidence'] for r in rs if r['Closure Evidence']})) or 'pending'} |
| **Reviewer comment** | {' '.join(sorted({r['Reviewer Comment'] for r in rs if r['Reviewer Comment']})) or '—'} |
""")
    flags = md_table([{"Flag": k, "Meaning": v["meaning"], "Rows": v["count"]} for k, v in summary["flags_by_type"].items()],
                     ["Flag", "Meaning", "Rows"])
    by_find = md_table([{"Finding": k, "Title": D["fnd"][k]["title"], "Exceptions": v}
                        for k, v in sorted(summary["exceptions_by_finding"].items())], ["Finding", "Title", "Exceptions"])
    txt = f"""# Access-Control Review: Findings

{GEN}

Full certification: [access-review.csv](access-review.csv) · [access-review.xlsx](access-review.xlsx) · raw exports: [source-exports/](source-exports/) · engine: [tools/review_engine.py](../tools/review_engine.py) · process: [README.md](README.md)

## Results at a glance
| Measure | Value |
|---|---|
| Snapshot date | {summary['snapshot_date']} |
| Entitlements reviewed | **{summary['entitlement_rows_reviewed']}** ({summary['privileged_rows']} privileged · {summary['phi_rows']} PHI-bearing) |
| Accounts reviewed | {summary['identities_reviewed']} across {len(summary['systems'])} system labels |
| Rows flagged by engine | {summary['flagged_rows']} |
| **Confirmed exceptions** | **{summary['exceptions']}** across {summary['exception_rows']} rows ({summary['exception_rate_pct']}% of rows) |
| False positives dispositioned | {summary['false_positive_rows']} (IR-Responder role inactive by design) |
| Exceptions closed · risk-accepted · open | {summary['exceptions_closed']} · {summary['exceptions_risk_accepted']} · {summary['exceptions_open']} |
| GitHub org owners (max {summary['github_owner_max']}) | {summary['github_owner_count']} at snapshot |

Engine flags are **not** findings. A reviewer confirms or dismisses each one. Exceptions are then **rolled up by root cause** into control-level findings, because auditors report control failures, not individual accounts.

## Exceptions rolled up to findings
{by_find}

## Engine flags
{flags}

## Exception index
{md_table(table, ['Exception', 'Accounts', 'Decision', 'Finding', 'Status'])}

## Exception detail: FINDING → RISK → CONTROL → REMEDIATION → OWNER → DUE → EVIDENCE OF CLOSURE
{chr(10).join(blocks)}
"""
    (ROOT / "access-review" / "findings.md").write_text(txt, encoding="utf-8")


# ------------------------------------------------------------------ evidence
EV_HEADERS = ["Evidence ID", "Control ID", "Evidence Request", "Artifact Description", "Evidence Owner", "Request Date",
              "Due Date", "Received Date", "Status", "Evidence Period", "File / Repository Location", "In Repo?",
              "Reviewer", "Review Result", "Evidence Level", "Sufficiency", "Exception", "Follow-Up Required",
              "Days Late", "Notes"]


def evidence_rows(D):
    return [{"Evidence ID": e["id"], "Control ID": "; ".join(e["controls"]), "Evidence Request": e["request"],
             "Artifact Description": e["artifact"], "Evidence Owner": e["owner"], "Request Date": d(e["request_date"]),
             "Due Date": d(e["due_date"]), "Received Date": d(e.get("received")), "Status": e["status"],
             "Evidence Period": e["period"], "File / Repository Location": e["location"], "In Repo?": e["in_repo"],
             "Reviewer": e["reviewer"], "Review Result": e["review_result"], "Evidence Level": e["level"],
             "Sufficiency": e["sufficiency"], "Exception": e["exception"], "Follow-Up Required": e["follow_up"],
             "Days Late": e["days_late"], "Notes": e.get("notes", "")} for e in D["evidence"]]


def sheet_evidence(wb, D):
    rows = evidence_rows(D)
    ws = table_sheet(wb, "Evidence", EV_HEADERS, [[r[h] for h in EV_HEADERS] for r in rows],
                     {"Evidence Request": 42, "Artifact Description": 40, "Review Result": 44, "File / Repository Location": 44,
                      "Notes": 36, "Exception": 28, "Follow-Up Required": 28},
                     {"Evidence Request", "Artifact Description", "Review Result", "Notes", "Exception", "Follow-Up Required",
                      "File / Repository Location"}, {"Request Date", "Due Date", "Received Date"})
    rec, due, late = (col_letter(EV_HEADERS, x) for x in ("Received Date", "Due Date", "Days Late"))
    for r in range(2, ws.max_row + 1):
        ws[f"{late}{r}"] = (f'=IF({rec}{r}="",MAX(0,{ASOF}-{due}{r}),MAX(0,{rec}{r}-{due}{r}))')
    add_list_validation(ws, EV_HEADERS, "Status", VOCAB["evidence_status"])
    add_list_validation(ws, EV_HEADERS, "Evidence Level", VOCAB["level"])
    status_formatting(ws, EV_HEADERS, "Status")
    return rows


def build_evidence(D):
    wb = new_wb()
    readme_sheet(wb, "Audit-Evidence Tracker (PBC)", [
        ("Statuses", ", ".join(VOCAB["evidence_status"])),
        ("Evidence level", "Design = documentation exists; Implementation = configured/deployed; Operating Effectiveness = worked over the period. See audit/evidence-quality-guide.md"),
        ("Sufficiency", "Derived: Accepted/Closed = Sufficient; Insufficient/Rework = Insufficient; others Pending"),
        ("Days Late", "Live formula: received after due, or still outstanding past due as of the README date"),
        ("Completion %", "(Accepted + Closed) / all requested items. 'Insufficient' is NOT complete: it becomes a finding"),
        ("Locations", "Paths starting with evidence/ or access-review/ are synthetic samples committed in this repo; 'GRC-Vault:/' paths are the fictional evidence vault"),
    ], STATUS_DATE)
    rows = sheet_evidence(wb, D)
    wb.save(ROOT / "audit" / "evidence-tracker.xlsx")
    write_csv(ROOT / "audit" / "evidence-tracker.csv", rows, EV_HEADERS)
    st = Counter(e["status"] for e in D["evidence"])
    lvl = Counter(e["level"] for e in D["evidence"])
    fam = defaultdict(list)
    for e in D["evidence"]:
        fam[e["controls"][0].split("-")[0]].append(e)
    sections = []
    for k in sorted(fam):
        sections.append(f"### {k}: {FAMILY[k]}\n" + md_table(
            [{"ID": e["id"], "Controls": ", ".join(e["controls"]), "Request (what, period, format)": f"{e['request']} · *{e['period']}*",
              "Owner": e["owner"], "Due": e["due_date"], "Status": e["status"], "Level": e["level"],
              "Sample in repo": f"[file](../{e['location']})" if e["in_repo"] == "Yes" else "—"} for e in fam[k]],
            ["ID", "Controls", "Request (what, period, format)", "Owner", "Due", "Status", "Level", "Sample in repo"]))
    total = len(D["evidence"])
    complete = st["Accepted"] + st["Closed"]
    txt = f"""# Evidence Request List (PBC) and Status

{GEN}

Tracker: [evidence-tracker.xlsx](evidence-tracker.xlsx) · [CSV](evidence-tracker.csv) · Quality rules: [evidence-quality-guide.md](evidence-quality-guide.md)

**How requests were written.** Each request names the **artifact**, the **period** or as-of date, the **population** ("ALL terminations", not "some examples"), and the **format** (system-generated export over screenshot). Vague requests produce vague evidence.

## Status as of {STATUS_DATE}
| Status | Count |
|---|---|
{chr(10).join(f'| {s} | {st[s]} |' for s in VOCAB['evidence_status'] if st[s])}
| **Total** | **{total}** |

**Evidence completion = (Accepted + Closed) / requested = {complete} / {total} = {round(100 * complete / total)}%.** Insufficient items are not "complete". Each one became part of a finding.

| Highest evidence level | Items |
|---|---|
{chr(10).join(f'| {k} | {lvl[k]} |' for k in VOCAB['level'])}

## Requests by control family
{chr(10).join(sections)}
"""
    (ROOT / "audit" / "evidence-request-list.md").write_text(txt, encoding="utf-8")
    return rows


# ------------------------------------------------------------------ findings
FIND_HEADERS = ["Finding ID", "Title", "Severity", "Likelihood", "Impact", "Score", "Primary Control", "Related Controls",
                "Related Risks", "POA&M ID", "Source", "Related Items (AR-EX / VR / EVID)", "Date Identified", "Owner",
                "Status", "Condition", "Criteria", "Cause", "Effect", "Recommendation", "Management Response",
                "Challenge Record"]


def finding_rows(D):
    return [{"Finding ID": f["id"], "Title": f["title"], "Severity": f["severity"], "Likelihood": f["likelihood"],
             "Impact": f["impact"], "Score": f["score"], "Primary Control": f["primary_control"],
             "Related Controls": "; ".join(f["controls"]), "Related Risks": "; ".join(f["risks"]), "POA&M ID": f["poam"],
             "Source": f["source"], "Related Items (AR-EX / VR / EVID)": "; ".join(f["related"]),
             "Date Identified": d(f["identified"]), "Owner": f["owner"], "Status": f["status"],
             "Condition": f["condition"], "Criteria": f["criteria"], "Cause": f["cause"], "Effect": f["effect"],
             "Recommendation": f["recommendation"], "Management Response": f["management_response"],
             "Challenge Record": f.get("challenge_record", "")} for f in D["findings"]]


def sheet_findings(wb, D):
    rows = finding_rows(D)
    wrap = {"Title", "Condition", "Criteria", "Cause", "Effect", "Recommendation", "Management Response",
            "Challenge Record", "Related Items (AR-EX / VR / EVID)"}
    ws = table_sheet(wb, "Findings", FIND_HEADERS, [[r[h] for h in FIND_HEADERS] for r in rows],
                     {h: 44 for h in wrap} | {"Title": 34}, wrap, {"Date Identified"})
    L, I, S, B = (col_letter(FIND_HEADERS, x) for x in ("Likelihood", "Impact", "Score", "Severity"))
    for r in range(2, ws.max_row + 1):
        ws[f"{S}{r}"] = f"={L}{r}*{I}{r}"
        ws[f"{B}{r}"] = rating_formula(f"{S}{r}")
    band_formatting(ws, FIND_HEADERS, "Severity")
    add_list_validation(ws, FIND_HEADERS, "Status", VOCAB["poam_status"])
    status_formatting(ws, FIND_HEADERS, "Status")
    return rows


def build_findings(D):
    wb = new_wb()
    readme_sheet(wb, "Findings Register", [
        ("Format", "Audit '5 Cs': Condition (what we found), Criteria (the requirement), Cause (root cause), Effect (risk), Recommendation"),
        ("Severity", "Live formula: band(Likelihood x Impact) at the level of the individual weakness (RSK-MTH-001 §10)"),
        ("Communication", "Each finding also has an executive version (findings/findings-register.md)"),
        ("Linkage", "FIND-nnn <-> POAM-nnn share numbers; every finding maps to >= 1 control and >= 1 enterprise risk"),
    ], STATUS_DATE)
    rows = sheet_findings(wb, D)
    obs = [[o["id"], o["control"], o["text"]] for o in D["obs"].values()]
    table_sheet(wb, "Observations", ["Observation ID", "Control", "Observation"], obs, {"Observation": 110}, {"Observation"})
    wb.save(ROOT / "findings" / "findings-register.xlsx")
    write_csv(ROOT / "findings" / "findings-register.csv", rows, FIND_HEADERS)
    sev = Counter(f["severity"] for f in D["findings"])
    stat = Counter(f["status"] for f in D["findings"])
    idx = md_table([{"ID": f"[{f['id']}](#{f['id'].lower()})", "Title": f["title"], "Severity": f["severity"],
                     "Control": f["primary_control"], "Risk": ", ".join(f["risks"]), "POA&M": f["poam"], "Status": f["status"]}
                    for f in D["findings"]], ["ID", "Title", "Severity", "Control", "Risk", "POA&M", "Status"])
    blocks = []
    for f in D["findings"]:
        ex = f["exec"]
        chal = f"\n**Challenge record (owner disputed):** {f['challenge_record']}\n" if f.get("challenge_record") else ""
        blocks.append(f"""<a id="{f['id'].lower()}"></a>
## {f['id']}: {f['title']}
**Severity:** {f['severity']} (L{f['likelihood']} × I{f['impact']} = {f['score']}) · **Status:** {f['status']} · **Owner:** {f['owner']} · **Identified:** {f['identified']} · **Source:** {f['source']}
**Controls:** {', '.join(f['controls'])} · **Risk:** {', '.join(f['risks'])} · **POA&M:** [{f['poam']}](../poam/poam.md#{f['poam'].lower()}) · **Related:** {', '.join(f['related'])}

### Technical / GRC version
| | |
|---|---|
| **Condition** | {f['condition']} |
| **Criteria** | {f['criteria']} |
| **Cause** | {f['cause']} |
| **Effect** | {f['effect']} |
| **Recommendation** | {f['recommendation']} |
| **Management response** | {f['management_response']} |
{chal}
### Executive version
- **What happened:** {ex['what_happened']}
- **Why it matters:** {ex['why_care']}
- **Business impact:** {ex['business_impact']}
- **Recommended action / status:** {ex['action']}
- **Urgency:** {ex['urgency']}
""")
    obs_md = md_table([{"ID": o["id"], "Control": o["control"], "Observation": o["text"]} for o in D["obs"].values()],
                      ["ID", "Control", "Observation"])
    txt = f"""# Findings Register

{GEN}

Workbook: [findings-register.xlsx](findings-register.xlsx) · CSV: [findings-register.csv](findings-register.csv) · POA&M: [poam.md](../poam/poam.md)

**{len(D['findings'])} findings**: {', '.join(f'{sev[s]} {s}' for s in VOCAB['severity'] if sev[s])} · Status as of {STATUS_DATE}: {', '.join(f'{v} {k}' for k, v in stat.items())}. Plus {len(D['obs'])} observations (below the finding threshold).

Each finding is written twice. The **Technical / GRC version** uses the audit "5 Cs" for control owners and auditors. The **Executive version** is plain language for leadership: what happened, why it matters, business impact, action, and urgency.

{idx}

{chr(10).join(blocks)}
## Observations
{obs_md}
"""
    (ROOT / "findings" / "findings-register.md").write_text(txt, encoding="utf-8")
    return rows


# ------------------------------------------------------------------ POA&M
POAM_HEADERS = ["POA&M ID", "Related Finding", "Related Control", "Related Risk", "Weakness", "Root Cause", "Severity",
                "Corrective Action", "Milestones", "Owner", "Resources Required", "Open Date", "SLA Date",
                "Original Target Date", "Current Target Date", "Target Change Reason", "Current Status",
                "Percent Complete", "Days Open", "Overdue?", "Slipped?", "Target Within SLA?", "Dependencies",
                "Validation Method", "Closure Evidence", "Validated By", "Closure Date", "Residual Risk", "Notes"]
MS_HEADERS = ["POA&M ID", "Milestone", "Description", "Due Date", "Status", "Completed Date", "Late?"]


def poam_rows(D):
    rows = []
    for p in D["poam"]:
        rows.append({"POA&M ID": p["id"], "Related Finding": p["finding"], "Related Control": "; ".join(p["controls"]),
                     "Related Risk": "; ".join(p["risks"]), "Weakness": p["weakness"], "Root Cause": p["root_cause"],
                     "Severity": p["severity"], "Corrective Action": p["corrective_action"],
                     "Milestones": " | ".join(f"{m['id']} {m['desc']} (due {m['due']}; {m['status']})" for m in p["milestones"]),
                     "Owner": p["owner"], "Resources Required": p["resources"], "Open Date": p["open"],
                     "SLA Date": p["sla_date"], "Original Target Date": d(p["original_target"]),
                     "Current Target Date": p["target"], "Target Change Reason": p.get("target_change_reason", ""),
                     "Current Status": p["status"], "Percent Complete": round(p["pct"], 2), "Days Open": p["days_open"],
                     "Overdue?": "Yes" if p["overdue"] else "No", "Slipped?": "Yes" if p["slipped"] else "No",
                     "Target Within SLA?": "n/a (acceptance expiry)" if p["status"] == "Risk Accepted" else ("Yes" if p["within_sla"] else "No"), "Dependencies": p["dependencies"],
                     "Validation Method": p["validation_method"], "Closure Evidence": "; ".join(p["closure_evidence"]),
                     "Validated By": p["validated_by"], "Closure Date": p["closed"], "Residual Risk": p["residual_risk"],
                     "Notes": p["notes"]})
    return rows


def milestone_rows(D):
    return [{"POA&M ID": p["id"], "Milestone": m["id"], "Description": m["desc"], "Due Date": d(m["due"]),
             "Status": m["status"], "Completed Date": d(m["completed"]),
             "Late?": "Yes" if (m["status"] != "Done" and d(m["due"]) < STATUS_DATE) else "No"}
            for p in D["poam"] for m in p["milestones"]]


def sheet_poam(wb, D):
    rows = poam_rows(D)
    wrap = {"Weakness", "Root Cause", "Corrective Action", "Milestones", "Resources Required", "Target Change Reason",
            "Dependencies", "Validation Method", "Residual Risk", "Notes"}
    ws = table_sheet(wb, "POAM", POAM_HEADERS, [[r[h] for h in POAM_HEADERS] for r in rows],
                     {h: 42 for h in wrap} | {"Milestones": 70}, wrap,
                     {"Open Date", "SLA Date", "Original Target Date", "Current Target Date", "Closure Date"}, {"Percent Complete"})
    c = {h: col_letter(POAM_HEADERS, h) for h in POAM_HEADERS}
    for r in range(2, ws.max_row + 1):
        ws[f"{c['SLA Date']}{r}"] = (f'={c["Open Date"]}{r}+IF({c["Severity"]}{r}="Critical",30,IF({c["Severity"]}{r}="High",90,'
                                     f'IF({c["Severity"]}{r}="Moderate",180,365)))')
        ws[f"{c['Percent Complete']}{r}"] = (f'=IFERROR(COUNTIFS(Milestones!$A:$A,{c["POA&M ID"]}{r},Milestones!$E:$E,"Done")'
                                             f'/COUNTIF(Milestones!$A:$A,{c["POA&M ID"]}{r}),0)')
        ws[f"{c['Days Open']}{r}"] = (f'=IF({c["Current Status"]}{r}="Closed",{c["Closure Date"]}{r}-{c["Open Date"]}{r},'
                                      f'{ASOF}-{c["Open Date"]}{r})')
        ws[f"{c['Overdue?']}{r}"] = (f'=IF(AND({c["Current Status"]}{r}<>"Closed",{c["Current Status"]}{r}<>"Risk Accepted",'
                                     f'{c["Current Target Date"]}{r}<{ASOF}),"Yes","No")')
        ws[f"{c['Slipped?']}{r}"] = f'=IF({c["Current Target Date"]}{r}<>{c["Original Target Date"]}{r},"Yes","No")'
        ws[f"{c['Target Within SLA?']}{r}"] = (f'=IF({c["Current Status"]}{r}="Risk Accepted","n/a (acceptance expiry)",'
                                              f'IF({c["Current Target Date"]}{r}<={c["SLA Date"]}{r},"Yes","No"))')
    add_list_validation(ws, POAM_HEADERS, "Current Status", VOCAB["poam_status"])
    band_formatting(ws, POAM_HEADERS, "Severity")
    status_formatting(ws, POAM_HEADERS, "Current Status")
    status_formatting(ws, POAM_HEADERS, "Overdue?")
    mrows = milestone_rows(D)
    ms = table_sheet(wb, "Milestones", MS_HEADERS, [[r[h] for h in MS_HEADERS] for r in mrows],
                     {"Description": 70}, {"Description"}, {"Due Date", "Completed Date"})
    for r in range(2, ms.max_row + 1):
        ms[f"G{r}"] = f'=IF(AND(E{r}<>"Done",D{r}<{ASOF}),"Yes","No")'
    add_list_validation(ms, MS_HEADERS, "Status", ["Not Started", "In Progress", "Delayed", "Done"])
    status_formatting(ms, MS_HEADERS, "Late?")
    return rows, mrows


def build_poam(D):
    wb = new_wb()
    readme_sheet(wb, "Plan of Action & Milestones (POA&M)", [
        ("Scope", "Every finding (FIND-nnn) has POAM-nnn, including the risk-accepted item: nothing silently drops off"),
        ("SLA", "From Open Date: Critical 30 / High 90 / Moderate 180 / Low 365 days (RSK-MTH-001 §9). SLA Date is a live formula"),
        ("Original vs current target", "Original target never changes. A changed current target needs a documented, approved reason; 'Slipped?' is a formula"),
        ("Percent Complete", "Formula: milestones Done / total milestones (Milestones sheet). No subjective percentages"),
        ("Overdue?", "Formula: status not Closed/Risk Accepted AND current target < as-of date"),
        ("Closure rule", "Status 'Closed' only after assessor validation with closure evidence (see Validation Method)"),
    ], STATUS_DATE)
    rows, mrows = sheet_poam(wb, D)
    wb.save(ROOT / "poam" / "poam.xlsx")
    write_csv(ROOT / "poam" / "poam.csv", rows, POAM_HEADERS)
    write_csv(ROOT / "poam" / "poam-milestones.csv", mrows, MS_HEADERS)
    st = Counter(p["status"] for p in D["poam"])
    idx = md_table([{"ID": f"[{p['id']}](#{p['id'].lower()})", "Finding": p["finding"], "Severity": p["severity"],
                     "Owner": p["owner"].split(" (")[0], "Open": p["open"], "Target": p["target"],
                     "Status": p["status"], "% done": f"{round(100 * p['pct'])}%",
                     "Flags": " ".join(x for x, y in (("⏰ OVERDUE", p["overdue"]), ("↪ SLIPPED", p["slipped"]),
                                                      ("⚠ beyond SLA", not p["within_sla"])) if y) or "—"}
                    for p in D["poam"]], ["ID", "Finding", "Severity", "Owner", "Open", "Target", "Status", "% done", "Flags"])
    blocks = []
    for p in D["poam"]:
        ms = md_table([{"#": m["id"], "Milestone": m["desc"], "Due": m["due"], "Status": m["status"],
                        "Completed": m["completed"] or "—"} for m in p["milestones"]], ["#", "Milestone", "Due", "Status", "Completed"])
        blocks.append(f"""<a id="{p['id'].lower()}"></a>
### {p['id']}: {p['weakness']}
**Finding** {p['finding']} · **Severity** {p['severity']} · **Controls** {', '.join(p['controls'])} · **Risks** {', '.join(p['risks'])} · **Owner** {p['owner']}
**Open** {p['open']} · **SLA** {p['sla_date']} · **Original target** {p['original_target']} · **Current target** {p['current_target']} · **Status** {p['status']} ({round(100 * p['pct'])}% of milestones)

- **Root cause:** {p['root_cause']}
- **Corrective action:** {p['corrective_action']}
- **Resources:** {p['resources']}
- **Dependencies:** {p['dependencies'] or '—'}
- **Validation method:** {p['validation_method']}
- **Closure evidence:** {', '.join(p['closure_evidence']) or '—'} · **Validated by:** {p['validated_by'] or '—'} · **Closed:** {p['closure_date'] or '—'}
- **Residual risk:** {p['residual_risk']}
{('- **Target change:** ' + p['target_change_reason']) if p.get('target_change_reason') else ''}
{('- **Notes:** ' + p['notes']) if p['notes'] else ''}

{ms}
""")
    txt = f"""# Plan of Action & Milestones (POA&M)

{GEN}

Workbook with live formulas (SLA date, % complete, days open, overdue, slipped): [poam.xlsx](poam.xlsx) · CSV: [poam.csv](poam.csv), [poam-milestones.csv](poam-milestones.csv) · Full lifecycle traces: [lifecycle-traces.md](lifecycle-traces.md)

**Status as of {STATUS_DATE}:** {', '.join(f'{v} {k}' for k, v in st.items())} · **Overdue:** {sum(p['overdue'] for p in D['poam'])} ({', '.join(p['id'] for p in D['poam'] if p['overdue']) or 'none'}) · **Slipped:** {sum(p['slipped'] for p in D['poam'])} ({', '.join(p['id'] for p in D['poam'] if p['slipped']) or 'none'}) · **Late milestones:** {sum(len(p['late_milestones']) for p in D['poam'])}

{idx}

## Items
{chr(10).join(blocks)}
"""
    (ROOT / "poam" / "poam.md").write_text(txt, encoding="utf-8")

    # lifecycle traces for every closed item
    traces = []
    for p in D["poam"]:
        if p["status"] != "Closed":
            continue
        f = D["fnd"][p["finding"]]
        ctl = D["ctl"][f["primary_control"]]
        risk = D["rsk"][f["risks"][0]]
        first_done = min((m for m in p["milestones"]), key=lambda m: str(m["due"]))
        traces.append(f"""## {f['id']} → {p['id']}: {f['title']}

| Stage | What happened | Date | Evidence |
|---|---|---|---|
| **1. Control failure** | {ctl['id']} {ctl['name']}: {ctl['test_summary'].split(' Validation')[0].split(' Post-remediation')[0]} | fieldwork | {', '.join(ctl['evidence'][:4])} |
| **2. Audit finding** | {f['id']} ({f['severity']}): {f['condition']} | {f['identified']} | {', '.join(x for x in f['related'] if x.startswith(('EVID', 'AR-EX', 'VR')))} |
| **3. Risk** | {risk['id']} {risk['title']}: current {risk['current_score']} {risk['current_band']} (L{risk['current'][0]}×I{risk['current'][1]}) | {risk['review_date']} | [risk register](../risk/risk-register.md) |
| **4. POA&M** | {p['id']} opened; owner {p['owner']}; target {p['current_target']} (SLA {p['sla_date']}) | {p['open_date']} | [poam.md](poam.md#{p['id'].lower()}) |
| **5. Corrective action** | {p['corrective_action']} ({len(p['milestones'])} milestones, first due {first_done['due']}) | {', '.join(str(m['completed']) for m in p['milestones'][:-1])} | — |
| **6. Validation** | {p['validation_method']} | {p['milestones'][-1]['completed']} | {', '.join(p['closure_evidence'])} |
| **7. Closure** | Closed by {p['validated_by']}; residual {p['residual_risk']} | {p['closure_date']} | — |
""")
    (ROOT / "poam" / "lifecycle-traces.md").write_text(
        f"""# Remediation Lifecycle Traces

{GEN}

The spec requires at least 5 findings to travel the **full** lifecycle. **{len(traces)} did**, each closed only after independent validation:

```
CONTROL FAILURE → AUDIT FINDING → RISK → POA&M → CORRECTIVE ACTION → VALIDATION → CLOSURE
```

Closure means the assessor **re-performed** the test (or observed the control operating) and inspected closure evidence. The owner saying "done" is not enough. Compare POAM-003, where a design-only screenshot (EVID-071) was **returned as Rework Required**.

{chr(10).join(traces)}""", encoding="utf-8")
    return rows


# ------------------------------------------------------------------ vendor
def vendor_inventory_rows(V):
    out = []
    for v in V["inventory"]:
        s = sum(v["scores"])
        tier = 1 if s >= 9 else 2 if s >= 5 else 3
        last = d(v["last_assessed"]) if v["last_assessed"] else None
        months = {1: 12, 2: 24, 3: 36}[tier]
        nxt = dt.date(last.year + (last.month - 1 + months) // 12, (last.month - 1 + months) % 12 + 1, min(last.day, 28)) if last else None
        current = "No (never)" if not last else ("No (overdue)" if nxt < STATUS_DATE else "Yes")
        out.append({"Vendor ID": v["id"], "Vendor": v["name"], "Service": v["service"], "Business Owner": v["owner"],
                    "Data Sensitivity (0-3)": v["scores"][0], "Data Volume (0-3)": v["scores"][1],
                    "Access / Integration (0-3)": v["scores"][2], "Business Criticality (0-3)": v["scores"][3],
                    "Substitutability (0-2)": v["scores"][4], "Inherent Score": s, "Tier": tier,
                    "Last Assessed": last, "Reassess (months)": months, "Next Due": nxt, "Assessment Current?": current,
                    "Assessment Method": v["method"], "Breach Notice Term": v["breach_notice_term"],
                    "Fourth Parties Mapped?": v["fourth_party_mapped"],
                    "Real Platform?": "Yes (no claims made about provider)" if v["real_platform"] else "No (fictional)"})
    return out


VINV_HEADERS = ["Vendor ID", "Vendor", "Service", "Business Owner", "Data Sensitivity (0-3)", "Data Volume (0-3)",
                "Access / Integration (0-3)", "Business Criticality (0-3)", "Substitutability (0-2)", "Inherent Score",
                "Tier", "Last Assessed", "Reassess (months)", "Next Due", "Assessment Current?", "Assessment Method",
                "Breach Notice Term", "Fourth Parties Mapped?", "Real Platform?"]
VA_HEADERS = ["Question ID", "Security Domain", "Requirement", "Vendor Response", "Evidence", "Assessor Evaluation",
              "Risk Identified", "Severity", "Recommendation", "Owner", "Status", "Vendor Finding"]
VR_HEADERS = ["VR ID", "Title", "Severity", "Party Responsible", "Questions", "Description", "Recommendation", "Owner",
              "Status", "Related Finding", "Escalation"]


def sheet_vendor_inventory(wb, V):
    rows = vendor_inventory_rows(V)
    ws = table_sheet(wb, "VendorInventory", VINV_HEADERS, [[r[h] for h in VINV_HEADERS] for r in rows],
                     {"Vendor": 30, "Service": 40, "Assessment Method": 34, "Breach Notice Term": 26}, {"Service", "Assessment Method"},
                     {"Last Assessed", "Next Due"})
    c = {h: col_letter(VINV_HEADERS, h) for h in VINV_HEADERS}
    for r in range(2, ws.max_row + 1):
        ws[f"{c['Inherent Score']}{r}"] = f"=SUM({c['Data Sensitivity (0-3)']}{r}:{c['Substitutability (0-2)']}{r})"
        ws[f"{c['Tier']}{r}"] = f"=IF({c['Inherent Score']}{r}>=9,1,IF({c['Inherent Score']}{r}>=5,2,3))"
        ws[f"{c['Reassess (months)']}{r}"] = f"=IF({c['Tier']}{r}=1,12,IF({c['Tier']}{r}=2,24,36))"
        L = f"{c['Last Assessed']}{r}"
        ws[f"{c['Next Due']}{r}"] = f'=IF({L}="","",DATE(YEAR({L}),MONTH({L})+{c["Reassess (months)"]}{r},MIN(DAY({L}),28)))'
        ws[f"{c['Assessment Current?']}{r}"] = (f'=IF({L}="","No (never)",IF({c["Next Due"]}{r}<{ASOF},"No (overdue)","Yes"))')
    return rows


def vendor_assessment_rows(V):
    vr = {x["id"]: x for x in V["vendor_findings"]}
    rows = []
    for q in V["questionnaire"]:
        f = vr.get(q["vr"]) if q["vr"] else None
        rows.append({"Question ID": q["id"], "Security Domain": q["domain"], "Requirement": q["requirement"],
                     "Vendor Response": q["response"], "Evidence": q["evidence"], "Assessor Evaluation": q["evaluation"],
                     "Risk Identified": (f["title"] if f else ("None" + (f" ({q['note']})" if q.get("note") else ""))),
                     "Severity": f["severity"] if f else "", "Recommendation": f["recommendation"] if f else "",
                     "Owner": f["owner"] if f else "", "Status": f["status"] if f else "Closed (satisfactory)",
                     "Vendor Finding": q["vr"]})
    return rows


def vr_rows(V):
    return [{"VR ID": x["id"], "Title": x["title"], "Severity": x["severity"], "Party Responsible": x["party"],
             "Questions": ", ".join(x["questions"]), "Description": x["description"], "Recommendation": x["recommendation"],
             "Owner": x["owner"], "Status": x["status"], "Related Finding": x["finding"], "Escalation": x.get("escalated", "")}
            for x in V["vendor_findings"]]


def build_vendor(D):
    V = D["vendor"]
    # questionnaire (template + responses)
    wb = new_wb()
    readme_sheet(wb, "Third-Party Security Questionnaire", [
        ("Use", "Template sheet: send to vendors (fill yellow columns). Responses sheet: Quarrystone Analytics (fictional), assessed 2026-06-22 to 2026-07-24"),
        ("Evidence rule", "Self-attestation is not evidence. Each 'Yes' needs an artifact (report, config, policy excerpt). VR-001 was found by interview, not questionnaire"),
        ("Evaluation scale", ", ".join(VOCAB["vq_eval"])),
        ("Example row", "Row 2 of the Template sheet shows the expected answer format"),
    ], STATUS_DATE)
    tmpl = [[q["id"], q["domain"], q["requirement"], "", ""] for q in V["questionnaire"]]
    tmpl[0][3] = "EXAMPLE: CISO appointed 2022; policy set reviewed 2026-01; board update quarterly"
    tmpl[0][4] = "EXAMPLE: policy index (PDF), org chart"
    ws = table_sheet(wb, "Template", ["Question ID", "Security Domain", "Requirement", "Vendor Response", "Evidence Provided"],
                     tmpl, {"Requirement": 60, "Vendor Response": 50, "Evidence Provided": 40},
                     {"Requirement", "Vendor Response", "Evidence Provided"})
    from xlsx_style import INPUT_FILL
    for r in range(2, ws.max_row + 1):
        ws[f"D{r}"].fill = INPUT_FILL
        ws[f"E{r}"].fill = INPUT_FILL
    resp = [[q["id"], q["domain"], q["requirement"], q["response"], q["evidence"], q["evaluation"], q["vr"]] for q in V["questionnaire"]]
    ws2 = table_sheet(wb, "Quarrystone Responses", ["Question ID", "Security Domain", "Requirement", "Vendor Response",
                                                    "Evidence", "Assessor Evaluation", "Vendor Finding"], resp,
                      {"Requirement": 50, "Vendor Response": 55, "Evidence": 34}, {"Requirement", "Vendor Response", "Evidence"})
    add_list_validation(ws2, ["Question ID", "Security Domain", "Requirement", "Vendor Response", "Evidence",
                              "Assessor Evaluation", "Vendor Finding"], "Assessor Evaluation", VOCAB["vq_eval"])
    wb.save(ROOT / "vendor-risk" / "vendor-questionnaire.xlsx")

    # assessment workbook
    wb = new_wb()
    tiers = V["tiering"]
    readme_sheet(wb, "Vendor Risk Assessment: Quarrystone Analytics (fictional) + Vendor Inventory", [
        ("Tiering", "Inherent score = sensitivity(0-3) + volume(0-3) + access(0-3) + criticality(0-3) + substitutability(0-2); "
                    "Tier 1 >= 9, Tier 2 >= 5, else Tier 3. Formulas live in VendorInventory"),
        ("Reassessment", "; ".join(f"Tier {t['tier']}: {t['reassess']}" for t in tiers["tiers"])),
        ("Real platforms", "AWS, Okta, GitHub, Google Workspace appear only as inventory entries; no statement is made about their controls"),
        ("Decision", V["decision"]["treatment"] + ", approved " + V["decision"]["approved_by"]),
    ], STATUS_DATE)
    inv = sheet_vendor_inventory(wb, V)
    prof = [[k.replace("_", " ").title(), (v if not isinstance(v, list) else "")] for k, v in V["profile"].items() if k != "fourth_parties"]
    table_sheet(wb, "Profile", ["Attribute", "Value"], prof, {"Attribute": 24, "Value": 120}, {"Value"})
    fp = [[x["name"], x["role"], x["location"], x["in_soc2"], x["disclosed_initially"]] for x in V["profile"]["fourth_parties"]]
    table_sheet(wb, "FourthParties", ["Subcontractor (fourth party)", "Role", "Location", "In SOC 2 Report?", "Disclosed Initially?"],
                fp, {"Subcontractor (fourth party)": 40, "Role": 50}, {"Role"})
    va = vendor_assessment_rows(V)
    ws = table_sheet(wb, "VendorAssessment", VA_HEADERS, [[r[h] for h in VA_HEADERS] for r in va],
                     {"Requirement": 44, "Vendor Response": 48, "Risk Identified": 40, "Recommendation": 48, "Evidence": 28},
                     {"Requirement", "Vendor Response", "Risk Identified", "Recommendation", "Evidence", "Status"})
    band_formatting(ws, VA_HEADERS, "Severity")
    add_list_validation(ws, VA_HEADERS, "Assessor Evaluation", VOCAB["vq_eval"])
    vr = vr_rows(V)
    ws = table_sheet(wb, "VendorFindings", VR_HEADERS, [[r[h] for h in VR_HEADERS] for r in vr],
                     {"Title": 40, "Description": 60, "Recommendation": 50, "Status": 34, "Escalation": 34},
                     {"Title", "Description", "Recommendation", "Status", "Escalation"})
    band_formatting(ws, VR_HEADERS, "Severity")
    dec = V["decision"]
    dec_rows = [["Inherent vendor risk", dec["inherent"]], ["Material findings", ", ".join(dec["material_findings"])],
                ["Mitigating controls", dec["mitigating_controls"]], ["Residual risk", dec["residual"]],
                ["Recommended treatment", dec["treatment"]], ["Rationale", dec["rationale"]], ["Approved by", dec["approved_by"]],
                ["Conditions precedent (SHCA data)", " | ".join(dec["conditions_precedent_for_shca_data"])],
                ["Contract / security requirements", " | ".join(dec["contract_requirements"])],
                ["Monitoring requirements", " | ".join(dec["monitoring"])], ["Reassessment frequency", dec["reassessment"]]]
    table_sheet(wb, "RiskDecision", ["Element", "Decision"], dec_rows, {"Element": 30, "Decision": 130}, {"Decision"})
    wb.save(ROOT / "vendor-risk" / "vendor-assessment.xlsx")

    # markdown
    sev = Counter(x["severity"] for x in V["vendor_findings"])
    evc = Counter(q["evaluation"] for q in V["questionnaire"])
    inv_md = md_table([{"ID": r["Vendor ID"], "Vendor": r["Vendor"], "Service": r["Service"], "Score": r["Inherent Score"],
                        "Tier": r["Tier"], "Last assessed": r["Last Assessed"] or "never", "Current?": r["Assessment Current?"],
                        "Breach notice": r["Breach Notice Term"]} for r in inv],
                      ["ID", "Vendor", "Service", "Score", "Tier", "Last assessed", "Current?", "Breach notice"])
    t12 = [r for r in inv if r["Tier"] in (1, 2)]
    cur = sum(r["Assessment Current?"] == "Yes" for r in t12)
    va_md = md_table([{"ID": r["Question ID"], "Domain": r["Security Domain"], "Requirement": r["Requirement"],
                       "Vendor response": r["Vendor Response"], "Evaluation": r["Assessor Evaluation"],
                       "VR": r["Vendor Finding"] or "—"} for r in va],
                     ["ID", "Domain", "Requirement", "Vendor response", "Evaluation", "VR"])
    vr_md = md_table([{"ID": r["VR ID"], "Title": r["Title"], "Severity": r["Severity"], "Party": r["Party Responsible"],
                       "Status": r["Status"], "Finding": r["Related Finding"] or "monitor"} for r in vr],
                     ["ID", "Title", "Severity", "Party", "Status", "Finding"])
    fp_md = md_table([{"Subcontractor": x["name"], "Role": x["role"], "Location": x["location"], "In SOC 2?": x["in_soc2"],
                       "Disclosed initially?": x["disclosed_initially"]} for x in V["profile"]["fourth_parties"]],
                     ["Subcontractor", "Role", "Location", "In SOC 2?", "Disclosed initially?"])
    prof_md = "\n".join(f"| {k.replace('_', ' ').title()} | {v} |" for k, v in V["profile"].items()
                        if k not in ("fourth_parties", "vendor"))
    txt = f"""# Vendor Risk Assessment: Quarrystone Analytics (fictional)

{GEN}

Workbooks: [vendor-assessment.xlsx](vendor-assessment.xlsx) (inventory with tiering formulas, profile, fourth parties, assessment, findings, decision) · [vendor-questionnaire.xlsx](vendor-questionnaire.xlsx) (blank template + responses) · Executive summary: [vendor-summary.md](vendor-summary.md)

## 1. Program view: vendor inventory and tiering
Tier 1–2 vendors with a current assessment: **{cur} of {len(t12)}** (status date {STATUS_DATE}). This is the SR-6 test behind FIND-015.

{inv_md}

## 2. Vendor profile (VEN-004)
| Attribute | Detail |
|---|---|
{prof_md}

### Fourth-party (subcontractor) map
{fp_md}

## 3. Assessment ({len(va)} questions)
Evaluation mix: {', '.join(f'{v} {k}' for k, v in evc.items())}.

{va_md}

## 4. Vendor findings (VR-###)
{', '.join(f'{sev[s]} {s}' for s in VOCAB['severity'] if sev[s])}.

{vr_md}
"""
    (ROOT / "vendor-risk" / "vendor-assessment.md").write_text(txt, encoding="utf-8")
    return inv, va, vr


# ------------------------------------------------------------------ templates
TEMPLATES = {
    "control-matrix-template.csv": (CTRL_HEADERS, None),
    "risk-register-template.csv": (RISK_HEADERS, None),
    "access-review-template.csv": (AR_ORDER, None),
    "evidence-tracker-template.csv": (EV_HEADERS, None),
    "findings-register-template.csv": (FIND_HEADERS, None),
    "poam-template.csv": (POAM_HEADERS, None),
    "vendor-assessment-template.csv": (VA_HEADERS, None),
    "vendor-inventory-template.csv": (VINV_HEADERS, None),
}


def build_templates(D, examples):
    wb = new_wb()
    readme_sheet(wb, "GRC Templates (blank, with one example row each)", [
        ("How to use", "Each sheet is a blank template with the same columns as the populated artifact. Row 2 is an EXAMPLE drawn from the Wrenfield sample; delete it before use"),
        ("Controlled values", "Use the vocabularies documented in docs/id-scheme.md and docs/assessment/nist-800-53-strategy.md §5"),
        ("Populated examples", "See controls/, risk/, access-review/, audit/, findings/, poam/, vendor-risk/"),
    ], STATUS_DATE)
    for fname, (headers, _) in TEMPLATES.items():
        ex = examples.get(fname)
        write_csv(ROOT / "templates" / fname, [ex] if ex else [], headers)
        table_sheet(wb, fname.replace("-template.csv", "")[:31], headers, [[ex.get(h) for h in headers]] if ex else [])
    wb.save(ROOT / "templates" / "grc-templates.xlsx")


# ------------------------------------------------------------------ combined workbook + dashboard formulas
def build_combined(D):
    wb = new_wb()
    readme_sheet(wb, "Wrenfield GRC Workbook: all artifacts, one governance system", [
        ("Sheets", "Dashboard (live formulas) · Controls · Risks · Findings · POAM · Milestones · Evidence · AccessReview · VendorInventory · VendorAssessment · VendorFindings · Lists"),
        ("Why one workbook", "Cross-sheet formulas prove the artifacts share IDs: the Dashboard counts across sheets, so a status change anywhere updates management reporting"),
        ("Recalculation", "Built with openpyxl; Excel/Sheets/LibreOffice compute formulas on open (fullCalcOnLoad). Formula results were verified with pycel (tools/check_formulas.py)"),
        ("Source of truth", "data/*.yaml -> tools/build.py"),
    ], STATUS_DATE)
    dash = wb.create_sheet("Dashboard")
    sheet_controls(wb, D)
    sheet_risks(wb, D)
    sheet_findings(wb, D)
    sheet_poam(wb, D)
    sheet_evidence(wb, D)
    sheet_access(wb, D)
    V = D["vendor"]
    sheet_vendor_inventory(wb, V)
    va = vendor_assessment_rows(V)
    table_sheet(wb, "VendorAssessment", VA_HEADERS, [[r[h] for h in VA_HEADERS] for r in va])
    vr = vr_rows(V)
    table_sheet(wb, "VendorFindings", VR_HEADERS, [[r[h] for h in VR_HEADERS] for r in vr])
    lists = wb.create_sheet("Lists")
    for i, (k, vals) in enumerate(VOCAB.items(), start=1):
        lists.cell(row=1, column=i, value=k)
        for j, v in enumerate(vals, start=2):
            lists.cell(row=j, column=i, value=v)

    ctl, fin, pom, evi, acc, vfi, vin, rsk = ("Controls", "Findings", "POAM", "Evidence", "AccessReview", "VendorFindings",
                                              "VendorInventory", "Risks")
    col = lambda headers, name: col_letter(headers, name)  # noqa: E731
    cres = f"{ctl}!${col(CTRL_HEADERS, 'Assessment Result')}:${col(CTRL_HEADERS, 'Assessment Result')}"
    cval = f"{ctl}!${col(CTRL_HEADERS, 'Implementation Status (Assessor-Validated)')}:${col(CTRL_HEADERS, 'Implementation Status (Assessor-Validated)')}"
    cchg = f"{ctl}!${col(CTRL_HEADERS, 'Status Changed by Assessor?')}:${col(CTRL_HEADERS, 'Status Changed by Assessor?')}"
    fsev = f"{fin}!${col(FIND_HEADERS, 'Severity')}:${col(FIND_HEADERS, 'Severity')}"
    fst = f"{fin}!${col(FIND_HEADERS, 'Status')}:${col(FIND_HEADERS, 'Status')}"
    pst = f"{pom}!${col(POAM_HEADERS, 'Current Status')}:${col(POAM_HEADERS, 'Current Status')}"
    pod = f"{pom}!${col(POAM_HEADERS, 'Overdue?')}:${col(POAM_HEADERS, 'Overdue?')}"
    psl = f"{pom}!${col(POAM_HEADERS, 'Slipped?')}:${col(POAM_HEADERS, 'Slipped?')}"
    est = f"{evi}!${col(EV_HEADERS, 'Status')}:${col(EV_HEADERS, 'Status')}"
    eid = f"{evi}!$A:$A"
    aex = f"{acc}!${col(AR_ORDER, 'Exception ID')}:${col(AR_ORDER, 'Exception ID')}"
    vsev = f"{vfi}!${col(VR_HEADERS, 'Severity')}:${col(VR_HEADERS, 'Severity')}"
    vst = f"{vfi}!${col(VR_HEADERS, 'Status')}:${col(VR_HEADERS, 'Status')}"
    vtier = f"{vin}!${col(VINV_HEADERS, 'Tier')}:${col(VINV_HEADERS, 'Tier')}"
    vcur = f"{vin}!${col(VINV_HEADERS, 'Assessment Current?')}:${col(VINV_HEADERS, 'Assessment Current?')}"
    rcur = f"{rsk}!${col(RISK_HEADERS, 'Current Risk')}:${col(RISK_HEADERS, 'Current Risk')}"
    rres = f"{rsk}!${col(RISK_HEADERS, 'Residual Risk')}:${col(RISK_HEADERS, 'Residual Risk')}"
    metrics = [
        ("MET-01", "Controls in scope", f'=COUNTIF({ctl}!$A$2:$A$500,"?*")', "Count of control rows"),
        ("MET-02", "Controls assessed (excl. N/A)", f'=COUNTIF({cres},"Satisfied")+COUNTIF({cres},"Other Than Satisfied")', "Result is Satisfied or Other Than Satisfied"),
        ("MET-03", "Controls Satisfied", f'=COUNTIF({cres},"Satisfied")', "Result = Satisfied"),
        ("MET-03a", "% of assessed controls Satisfied", "=ROUND({MET-03}/{MET-02},2)", "MET-03 / MET-02"),
        ("MET-04a", "Validated: Implemented", f'=COUNTIF({cval},"Implemented")', "Assessor-validated status"),
        ("MET-04b", "Validated: Partially Implemented", f'=COUNTIF({cval},"Partially Implemented")', "Assessor-validated status"),
        ("MET-04c", "Validated: Planned or Not Implemented", f'=COUNTIF({cval},"Planned")+COUNTIF({cval},"Not Implemented")', "Assessor-validated status"),
        ("MET-05", "Controls where assessor changed owner-stated status", f'=COUNTIF({cchg},"Yes")', "Owner overstatement indicator"),
        ("MET-06", "Findings (total)", f'=COUNTIF({fin}!$A$2:$A$500,"FIND-*")', "All findings"),
        ("MET-06a", "Open findings (not Closed / Risk Accepted)", f'={{MET-06}}-COUNTIF({fst},"Closed")-COUNTIF({fst},"Risk Accepted")', "Total minus Closed minus Risk Accepted"),
        ("MET-07", "Open High/Critical findings", f'=COUNTIFS({fsev},"High",{fst},"<>Closed")+COUNTIFS({fsev},"Critical",{fst},"<>Closed")', "Severity High/Critical and not Closed"),
        ("MET-08", "Overdue POA&M items", f'=COUNTIF({pod},"Yes")', "Formula-driven Overdue? column"),
        ("MET-08a", "Slipped POA&M items (target changed)", f'=COUNTIF({psl},"Yes")', "Current target <> original target"),
        ("MET-09", "POA&M closed and validated", f'=COUNTIF({pst},"Closed")', "Status Closed (requires validation)"),
        ("MET-10", "Access-review exceptions (distinct)", f"=SUM({acc}!${col(AR_ORDER + ['First Row of Exception?'], 'First Row of Exception?')}:${col(AR_ORDER + ['First Row of Exception?'], 'First Row of Exception?')})", "Distinct AR-EX IDs (helper column counts each ID once)"),
        ("MET-10a", "Access-review exception rows", f'=COUNTIF({aex},"AR-EX-*")', "Rows carrying an exception"),
        ("MET-11", "Vendor findings (VR) total", f'=COUNTIF({vfi}!$A$2:$A$500,"VR-*")', "All VR items"),
        ("MET-11a", "Vendor findings Critical/High", f'=COUNTIF({vsev},"Critical")+COUNTIF({vsev},"High")', "Severity"),
        ("MET-11b", "Tier 1-2 vendors with current assessment", f'=COUNTIFS({vtier},"<=2",{vcur},"Yes")', "Formula-driven currency"),
        ("MET-11c", "Tier 1-2 vendors total", f'=COUNTIF({vtier},"<=2")', "Tier from inherent score"),
        ("MET-12", "Evidence completion %", f'=ROUND((COUNTIF({est},"Accepted")+COUNTIF({est},"Closed"))/COUNTIF({eid},"EVID-*"),2)', "(Accepted + Closed) / requested"),
        ("MET-13a", "Enterprise risks rated High+ (current)", f'=COUNTIF({rcur},"High")+COUNTIF({rcur},"Critical")', "Current (as tested)"),
        ("MET-13b", "Enterprise risks rated High+ (residual target)", f'=COUNTIF({rres},"High")+COUNTIF({rres},"Critical")', "Residual target"),
    ]
    dash["A1"] = "Executive Dashboard (live formulas)"
    dash["A2"] = "As of"
    dash["B2"] = f"={ASOF}"
    dash["B2"].number_format = "yyyy-mm-dd"
    dash.append([])
    dash.append(["Metric ID", "Metric", "Value", "Definition"])
    rowof = {m[0]: 5 + i for i, m in enumerate(metrics)}
    for mid, name, formula, defi in metrics:
        for k, r in rowof.items():
            formula = formula.replace("{" + k + "}", f"C{r}")
        dash.append([mid, name, formula, defi])
    from openpyxl.styles import Font
    from xlsx_style import FONT, HEADER_FILL
    dash["A1"].font = Font(name=FONT, bold=True, size=14, color="1F4E5A")
    for c in dash[4]:
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = HEADER_FILL
    for row in dash.iter_rows(min_row=5):
        for c in row:
            c.font = Font(name=FONT, size=10)
    for k in ("MET-03a", "MET-12"):
        dash[f"C{rowof[k]}"].number_format = "0%"
    dash.column_dimensions["A"].width = 10
    dash.column_dimensions["B"].width = 52
    dash.column_dimensions["C"].width = 12
    dash.column_dimensions["D"].width = 60
    (ROOT / "workbook").mkdir(exist_ok=True)
    wb.save(ROOT / "workbook" / "wrenfield-grc-workbook.xlsx")
    return [(m[0], m[1]) for m in metrics]


def main():
    D = load_all()
    ctl_rows = build_controls(D)
    build_worksheets_md(D)
    build_traceability(D)
    r_rows = build_risk(D)
    build_access_review(D)
    e_rows = build_evidence(D)
    f_rows = build_findings(D)
    p_rows = build_poam(D)
    inv, va, vr = build_vendor(D)
    examples = {
        "control-matrix-template.csv": next(r for r in ctl_rows if r["Control ID"] == "AC-2(3)"),
        "risk-register-template.csv": next(r for r in r_rows if r["Risk ID"] == "RISK-003"),
        "access-review-template.csv": next(r for r in D["ar"] if r["Row ID"] == "ENT-058"),
        "evidence-tracker-template.csv": next(r for r in e_rows if r["Evidence ID"] == "EVID-003"),
        "findings-register-template.csv": next(r for r in f_rows if r["Finding ID"] == "FIND-002"),
        "poam-template.csv": next(r for r in p_rows if r["POA&M ID"] == "POAM-013"),
        "vendor-assessment-template.csv": next(r for r in va if r["Question ID"] == "VQ-16"),
        "vendor-inventory-template.csv": next(r for r in inv if r["Vendor ID"] == "VEN-004"),
    }
    build_templates(D, examples)
    build_combined(D)
    print(f"built: {len(D['controls'])} controls, {len(D['risks'])} risks, {len(D['findings'])} findings, "
          f"{len(D['poam'])} POA&M, {len(D['evidence'])} evidence, {len(D['ar'])} access rows, "
          f"{len(D['vendor']['questionnaire'])} VQ, {len(D['vendor']['vendor_findings'])} VR")


if __name__ == "__main__":
    main()
